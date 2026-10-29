from pathlib import Path
import subprocess
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
ANALYSIS = ROOT / "stage2" / "report"
OUT = ROOT / "stage2" / "deliverables"
OUT.mkdir(exist_ok=True)

def chapter(path, number, title, offset=0, subprefix=None):
    text=path.read_text()
    lines=[]
    in_fence=False
    for line in text.splitlines():
        if line.lstrip().startswith('```'):
            in_fence=not in_fence
            lines.append(line)
            continue
        if in_fence:
            lines.append(line)
            continue
        if line.startswith("# "):
            if subprefix is not None:line=f"## {subprefix} {title}"
            else:line=f"# {number} {title}"
        elif re.match(r"#{2,4} ",line):
            match=re.match(r"(#{2,4})\s+(\d+(?:\.\d+)*)\.?\s*(.*)",line)
            if match:
                hashes,num,head=match.groups()
                if subprefix is not None:
                    line="#"+hashes+f" {subprefix}.{num} {head}"
                else:
                    nums=num.split(".");nums[0]=str(int(nums[0])+offset)
                    line=hashes+f" {number}."+".".join(nums)+f" {head}"
        lines.append(line)
    return "\n".join(lines)

scope=(ANALYSIS/"00_scope_versions.md").read_text()
v10=chapter(ROOT/"stage2/notes/v10_baseline.md",4,"原 v10 与当前整合稿逐项对照及旧基准复算")
v10=v10.replace("运行覆盖见第 8 节", "运行覆盖见第4.8节").replace("见第9节", "见第4.9节")
physical=(ANALYSIS/"01_physical_framework.md").read_text()
rod=chapter(ROOT/"stage2/notes/rod3d.md",6,"三维杆 血流 接触与磁驱动的实际计算")
rod += "\n\n## 6.10 计算结果图\n\n"
for filename,caption,width in [
    ("geometry_3d","图6-1 曲管内三维中心线的时间截面。各轴采用不同显示比例；该主跟踪例为小幅运动。",6.2),
    ("learned_feedforward_transfer","图6-2 同一12自由度前馈接口施加到32与64段杆。线性与精确低维ROM曲线接近，复网络在本例未优于它们。",6.85),
    ("full_coil_constraints","图6-3 从局部场分配到完整线圈及采样管域约束。原点分配违反实际梯度界；连续电流分支保留了可行性与性能代价。",6.85),
    ("refinement","图6-4 空间与时间细化。纵轴为相邻网格差，完整开放系统实测约一阶时间收敛。",6.85),
]:
    rod += "!["+caption+"]("+str(ROOT/"stage2/results/rod3d/figures"/(filename+".png"))+f"){{width={width}in}}\n\n"
bridge=(ANALYSIS/"03_birkhoff_bridge.md").read_text()
route=chapter(ROOT/"stage2/notes/route_a.md",7,"一般 Birkhoff 高维算法复算与误差审计",subprefix="7.7")
route=route.replace("第 7 节", "第7.7.7节").replace("第 6 节", "第7.7.6节")
learning=(ANALYSIS/"02_learning_metriplectic.md").read_text()
additional=[]
for filename,title,prefix in [("learning_confirmation.md","冻结模型后的新样本确认","8.6"),("learning_spd.md","机械正质量复网络的实际结果","8.7")]:
    path=ROOT/"stage2/notes"/filename
    if path.exists():additional.append(chapter(path,8,title,subprefix=prefix))
learning=learning.replace("# 9 Metriplectic", "\n\n"+"\n\n".join(additional)+"\n\n# 9 Metriplectic")
learning=learning.replace("# 9 Metriplectic", "\n\n![图8-1 冻结模型在新独立数据上的正则向量场误差。复网络显示三个种子的范围；正质量网络仅一个种子。容量不同，不据此推断复数表示优越。]("+str(ROOT/"stage2/results/figures/learning_confirmation.png")+"){width=6.85in}\n\n# 9 Metriplectic")
route += "\n\n![图7-1 16自由度合成变质量模型中的实测误差与收敛阶。各方法适用的二形式和阻尼条件见本节正文。]("+str(ROOT/"stage2/results/route_a/route_a_16dof_convergence.png")+"){width=6.85in}\n\n![图7-2 一般近似输运与Moser图中的两种误差。图匹配达到机器精度，并不意味着近似一形式准确表示原向量场。]("+str(ROOT/"stage2/results/route_a/route_a_moser_error_separation.png")+"){width=6.85in}"
end=(ANALYSIS/"04_synthesis_reproduction.md").read_text()
source="\n\n".join([scope,v10,physical,rod,bridge,route,learning,end])
def table_math_breaks(match):
    expression=match.group(1)
    if r"\begin" in expression:return match.group(0)
    expression=expression.replace(r"\left","").replace(r"\right","")
    parts=[];buffer=[];depth=0;i=0
    while i<len(expression):
        ch=expression[i]
        if ch=="{" and (i==0 or expression[i-1]!="\\"):depth+=1
        if ch=="}" and (i==0 or expression[i-1]!="\\"):depth-=1
        if depth==0 and expression.startswith(r"\quad",i):
            if buffer:parts.append("".join(buffer));buffer=[]
            i+=5;continue
        if depth==0 and expression.startswith(r"\times",i) and buffer:
            parts.append("".join(buffer));buffer=[r"\times"]
            i+=6;continue
        if depth==0 and ch in ["=","+",","] and buffer and not(i>0 and expression[i-1]=="_"):
            if ch==",":buffer.append(ch);parts.append("".join(buffer));buffer=[]
            else:parts.append("".join(buffer));buffer=[ch]
        else:buffer.append(ch)
        i+=1
    if buffer:parts.append("".join(buffer))
    return " ".join("$"+p.strip()+"$" for p in parts if p.strip())
source="\n".join(re.sub(r"(?<!\$)\$([^$]+)\$(?!\$)",table_math_breaks,line) if line.startswith("|") else line for line in source.splitlines())
source=re.sub(r"\$\$\s*\(4\.4576294,[\s\S]*?\$\$",lambda m:r"$$\begin{matrix}(4.4576294,4.4576294),\ (27.6941957,27.6941957),\\(79.7832866,79.7832866),\ (356.9402367,356.9402367)\ \mathrm{rad/s}.\end{matrix}$$",source)
# Make numbered formula annotations native Word math, avoiding unsupported tags.
source=re.sub(r"\\tag\{([^{}]+)\}",lambda m:r"\qquad\text{("+m.group(1)+")}",source)
source=source.replace(r"\!","")
source=re.sub(r"\\rm\s+([A-Za-z]+(?:\\\s+[A-Za-z]+)*)",lambda m:r"\mathrm{"+m.group(1)+"}",source)
source=source.replace(r"\bar",r"\overline")
source=source.replace(r"\begin{aligned}",r"\begin{matrix}").replace(r"\end{aligned}",r"\end{matrix}")
# The only align markers are equation alignment, not matrix columns.
source=source.replace("e_B&=", "e_B=").replace("\\\\\n&=", "\\\\\n=")
source=re.sub(r"\[([^\]]+)\]\(sandbox:[^)]+\)",r"\1",source)
source=source.replace("没有重新训练或改变候选环面。","证书复算没有重新训练或改变候选环面；独立的读出重训另列。")
md_path=ANALYSIS/"stage2_complete.md";md_path.write_text(source)
raw=ANALYSIS/"stage2_raw.docx"
log=ANALYSIS/"pandoc.log"
with log.open("w") as stream:
    subprocess.run(["pandoc",str(md_path),"-o",str(raw),"--from=markdown+tex_math_dollars-superscript-subscript","--standalone"],check=True,stderr=stream)
doc=Document(raw)
# LibreOffice's OMML importer substitutes a right parenthesis for some explicit
# square/absolute-value delimiters.  Keep these delimiters as ordinary native
# math runs, so interval endpoints and absolute values retain their meaning in
# both Word and the rendered preview.  None of these delimiters encloses a
# multirow matrix; their interior math objects remain editable and unchanged.
for delimiter in list(doc.element.xpath('.//m:d')):
    properties = delimiter.find(qn('m:dPr'))
    if properties is None:
        continue
    beginning = properties.find(qn('m:begChr'))
    ending = properties.find(qn('m:endChr'))
    begin_char = beginning.get(qn('m:val'), '(') if beginning is not None else '('
    end_char = ending.get(qn('m:val'), ')') if ending is not None else ')'
    if (begin_char, end_char) == ('(', ')'):
        continue
    arguments = delimiter.findall(qn('m:e'))
    if len(arguments) != 1:
        raise ValueError('Unexpected multi-argument non-round delimiter')
    parent = delimiter.getparent()
    position = parent.index(delimiter)
    replacement = []
    def delimiter_run(character):
        run = OxmlElement('m:r')
        run_properties = OxmlElement('m:rPr')
        style = OxmlElement('m:sty'); style.set(qn('m:val'), 'p')
        run_properties.append(style); run.append(run_properties)
        text = OxmlElement('m:t'); text.text = character; run.append(text)
        return run
    if begin_char:
        replacement.append(delimiter_run(begin_char))
    replacement.extend(list(arguments[0]))
    if end_char:
        replacement.append(delimiter_run(end_char))
    parent.remove(delimiter)
    for offset, element in enumerate(replacement):
        parent.insert(position + offset, element)
FONT = 'Noto Sans CJK SC'
styles_by_name = {style.name: style for style in doc.styles}

def set_font(style, size, bold=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    rp = style.element.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rp.insert(0, rf)
    for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + attr), FONT)
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + attr), None)

for style in doc.styles:
    if style.type == 1:
        set_font(style, 11)
        pf = style.paragraph_format
        pf.space_after = Pt(5)
        pf.line_spacing = 1.26
        pf.widow_control = True
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)

for name in ['Normal', 'Body Text', 'First Paragraph']:
    if name in styles_by_name:
        set_font(styles_by_name[name], 11)
        styles_by_name[name].paragraph_format.first_line_indent = Pt(0)
        styles_by_name[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

set_font(styles_by_name['Title'], 23, True)
styles_by_name['Title'].paragraph_format.space_after = Pt(12)
styles_by_name['Title'].paragraph_format.line_spacing = 1.15
set_font(styles_by_name['Subtitle'], 12, False)
styles_by_name['Subtitle'].paragraph_format.space_after = Pt(10)
for name, size in [('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)]:
    set_font(styles_by_name[name], size, True)
    pf = styles_by_name[name].paragraph_format
    pf.space_before = Pt(14 if name != 'Heading 1' else 18)
    pf.space_after = Pt(8)
    pf.keep_with_next = True
    pf.line_spacing = 1.15

for section in doc.sections:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.75)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.78)
    section.right_margin = Inches(.78)
    section.header_distance = Inches(.3)
    section.footer_distance = Inches(.3)
    hp = section.header.paragraphs[0]
    hp.text = '三维磁机器人第二阶段研究'
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.name = FONT; r.font.size = Pt(8); r.font.color.rgb = RGBColor(0,0,0)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run('第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    fp._p.append(field)
    fp.add_run(' 页')
    for r in fp.runs:
        r.font.name = FONT; r.font.size = Pt(9)

in_references = False
for pi, p in enumerate(doc.paragraphs):
    if p._p.xpath('.//w:rStyle[@w:val="VerbatimChar"]'):
        # Avoid large justified spaces before long paths and identifiers.
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if p.text == '参考文献与来源':
        in_references = True
    if in_references:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if p.style.name == 'Heading 1':
        p.paragraph_format.page_break_before = False
    if p._p.xpath('.//m:oMath'):
        p.paragraph_format.keep_together = True
        if p._p.xpath('./m:oMathPara'):
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
    if pi + 1 < len(doc.paragraphs):
        next_p = doc.paragraphs[pi + 1]
        if next_p._p.xpath('./m:oMathPara'):
            if len(p.text) < 65 or p._p.xpath('./m:oMathPara'):
                p.paragraph_format.keep_with_next = True
    for border in list(p._p.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)

# Deliberate widths for comparison tables, repeated header rows and non-splitting data rows.
total_width = 6.94
for table in doc.tables:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    heads = [c.text for c in table.rows[0].cells]
    n = len(heads)
    if False and heads[0].startswith('原编号'):
        fractions = [.10, .29, .14, .47]
    elif False and heads[0].startswith('编号及'):
        fractions = [.14, .38, .22, .26]
    elif False and heads[0] == '编号':
        fractions = [.07, .34, .42, .17]
    elif n == 3:
        fractions = [.24, .37, .39]
    elif n == 4:
        fractions = [.16, .30, .31, .23]
    elif n == 6:
        fractions = [.18, .22, .18, .10, .18, .14]
    elif n == 5:
        fractions = [.12, .21, .20, .27, .20]
    else:
        fractions = [1/n]*n
    for col, frac in zip(table.columns, fractions):
        col.width = Inches(total_width * frac)
    pr = table._tbl.tblPr
    borders = pr.find(qn('w:tblBorders'))
    if borders is None:
        borders = OxmlElement('w:tblBorders'); pr.append(borders)
    else:
        for child in list(borders): borders.remove(child)
    for side in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:'+side)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9')
        borders.append(el)
    for ri, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        if ri == 0:
            rep = OxmlElement('w:tblHeader'); trpr.append(rep)
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(total_width * fractions[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            for old in list(tcpr.findall(qn('w:tcBorders'))): tcpr.remove(old)
            cb = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el = OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); cb.append(el)
            tcpr.append(cb)
            mar = OxmlElement('w:tcMar')
            for side, v in [('top',75),('bottom',75),('left',85),('right',85)]:
                el = OxmlElement('w:'+side); el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa'); mar.append(el)
            tcpr.append(mar)
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'),'EAEAEA'); tcpr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.keep_with_next = False
                p.paragraph_format.keep_together = True
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if ci == 0 and fractions[ci] <= .16:
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = FONT; run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0,0,0)
                    if ri == 0:
                        run.bold = True
                    rp = run._r.get_or_add_rPr()
                    rf = rp.find(qn('w:rFonts'))
                    if rf is None:
                        rf = OxmlElement('w:rFonts'); rp.insert(0,rf)
                    rf.set(qn('w:eastAsia'),FONT)

# Math font and language defaults; keep source equations as OMML.
settings = doc.settings.element
mathPr = settings.find(qn('m:mathPr'))
if mathPr is None:
    mathPr = OxmlElement('m:mathPr'); settings.append(mathPr)
mathFont = mathPr.find(qn('m:mathFont'))
if mathFont is None:
    mathFont = OxmlElement('m:mathFont'); mathPr.append(mathFont)
mathFont.set(qn('m:val'),'Cambria Math')
for cp in doc.element.xpath('.//m:ctrlPr/w:rPr'):
    sz = cp.find(qn('w:sz'))
    if sz is None:
        sz = OxmlElement('w:sz'); cp.append(sz)
    sz.set(qn('w:val'),'22')

# A native inline equation does not wrap internally inside narrow table cells.
# The source splits long equalities at operators; use the table's own type scale.
for cp in doc.element.xpath('.//w:tbl//m:ctrlPr/w:rPr'):
    sz=cp.find(qn('w:sz'))
    if sz is None:sz=OxmlElement('w:sz');cp.append(sz)
    sz.set(qn('w:val'),'18')
for mr in doc.element.xpath('.//w:tbl//m:r'):
    rp=mr.find(qn('w:rPr'))
    if rp is None:rp=OxmlElement('w:rPr');mr.insert(0,rp)
    sz=rp.find(qn('w:sz'))
    if sz is None:sz=OxmlElement('w:sz');rp.append(sz)
    sz.set(qn('w:val'),'18')

dest = OUT / '1025_v10_第二阶段融合研究报告.docx'
doc.save(dest)
print(dest)
print(f'paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} math={len(doc.element.xpath(".//m:oMath"))}')
